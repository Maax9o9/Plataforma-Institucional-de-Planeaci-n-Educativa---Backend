"""Rellenar el formato POA de la SEP, celda por celda.

El `.xlsm` que la SEP entrega trae macros, rangos con nombre, celdas
combinadas y un formato oficial que no tiene sentido reproducir: lo que se
hace es **abrir una copia de la plantilla y escribir en sus celdas**. Lo que
sale es el mismo archivo que entró, con los datos de la plataforma dentro.

El formato apila *bloques* —una cédula por estrategia— dentro de la hoja de su
objetivo, y cada bloque tiene una capacidad fija: dos indicadores y cinco
actividades. No es una elección nuestra; es cómo viene el archivo. Por eso el
relleno nunca inserta filas: insertarlas rompería las celdas combinadas, el
área de impresión y las macros, y el resultado dejaría de ser el formato de la
SEP. Lo que no cabe se devuelve como incidencia, dicho con todas sus letras,
en vez de recortarse en silencio.

Las posiciones no están codificadas: se localizan por sus rótulos, igual que
hace el importador. Un formato que mueva una fila sigue funcionando; uno que
cambie los rótulos, no, y entonces hay que mirarlo.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any

#: Columnas del formato, por su letra. Se nombran para que el mapa de escritura
#: se lea como el Excel y no como una lista de números.
COL_B = 2
COL_C = 3
COL_D = 4
COL_E = 5
COL_J = 10
COL_M = 13
COL_P = 16
COL_R = 18
COL_S = 19
COL_V = 22
COL_X = 24
COL_Y = 25

#: Por cuatrimestre: (programado #, programado %, alcanzado #, alcanzado %).
#: F/G y H/I para el primero, J/K y L/M para el segundo, N/O y P/Q para el
#: tercero, tal como los rotula la fila de encabezado del calendario.
COLUMNAS_CUATRIMESTRE = {
    1: (6, 7, 8, 9),
    2: (10, 11, 12, 13),
    3: (14, 15, 16, 17),
}

CLAVE_ACTIVIDAD = re.compile(r"^\s*\d+\.\d+\.\d+\b")
CLAVE_ESTRATEGIA = re.compile(r"Estrategia\s+(\d+\.\d+)")
ANIO = re.compile(r"\b(20\d{2})\b")


@dataclass(frozen=True, kw_only=True)
class IndicadorDelFormato:
    name: str
    formula: str
    unit: str
    institutional_goal: Decimal | None = None
    baseline_year: int | None = None
    baseline_value: Decimal | None = None
    current_percentage: Decimal | None = None
    target_value: Decimal | None = None
    target_percentage: Decimal | None = None
    achieved_total: Decimal | None = None
    achieved_percentage: Decimal | None = None


@dataclass(frozen=True, kw_only=True)
class ActividadDelFormato:
    key: str
    description: str
    unit: str
    annual_goal: Decimal
    executing_area: str | None = None
    deviation_justification: str | None = None
    #: Cuatrimestre (1, 2, 3) → lo programado y lo alcanzado.
    scheduled: dict[int, Decimal | None] = field(default_factory=dict)
    achieved: dict[int, Decimal | None] = field(default_factory=dict)


@dataclass(frozen=True, kw_only=True)
class CedulaDelFormato:
    objective_number: int
    strategy_key: str
    strategy_denomination: str
    strategy_type: str | None = None
    responsible_area: str
    objective_denomination: str
    scope_and_socioeconomic_effect: str | None = None
    #: Los dos firmantes, como (nombre, cargo). El formato reserva dos huecos.
    signatories: tuple[tuple[str, str], ...] = ()
    indicators: tuple[IndicadorDelFormato, ...] = ()
    activities: tuple[ActividadDelFormato, ...] = ()


@dataclass
class Bloque:
    """Dónde escribe cada dato de una cédula dentro de su bloque."""

    hoja: str
    inicio: int
    estrategia: int
    area: int
    indicadores: list[int]
    objetivo: int
    alcance: int
    actividades_desde: int
    actividades_hasta: int

    @property
    def capacidad_actividades(self) -> int:
        return self.actividades_hasta - self.actividades_desde + 1

    @property
    def capacidad_indicadores(self) -> int:
        return len(self.indicadores)


def con_clave(clave: str, texto: str) -> str:
    """El texto precedido de su clave, sin repetirla.

    El catálogo del PSE guarda la descripción ya encabezada por la clave
    —«1.1.6 Desarrollar acciones…»— porque así viene en el documento federal,
    pero no siempre: la clave se antepone sólo cuando falta. Sin esta
    comprobación el formato salía diciendo «1.1.6 1.1.6 Desarrollar…».
    """
    limpio = texto.strip()
    if not limpio:
        return clave
    return limpio if limpio.startswith(clave) else f"{clave} {limpio}"


def _texto(valor: Any) -> str:
    return " ".join(str(valor).split()).lower() if isinstance(valor, str) else ""


def localizar_bloques(libro) -> list[Bloque]:
    """Los bloques rellenables del libro, en el orden en que aparecen.

    Sólo cuentan los que tienen calendario de actividades. El archivo trae
    además tres hojas del formato viejo —`académico`, `vinculación`,
    `gestión`— que hablan de «procesos» y no tienen calendario: son plantillas
    muertas y escribir en ellas produciría un documento que nadie pidió.
    """
    bloques: list[Bloque] = []
    for nombre in libro.sheetnames:
        hoja = libro[nombre]
        inicios = [
            fila
            for fila in range(1, hoja.max_row + 1)
            if "programa operativo anual" in _texto(hoja.cell(row=fila, column=COL_B).value)
        ]
        for indice, inicio in enumerate(inicios):
            fin = inicios[indice + 1] if indice + 1 < len(inicios) else hoja.max_row + 1
            bloque = _mapear(hoja, nombre, inicio, fin)
            if bloque is not None:
                bloques.append(bloque)
    return bloques


def _mapear(hoja, nombre: str, inicio: int, fin: int) -> Bloque | None:
    estrategia = area = objetivo = alcance = calendario = None
    indicadores: list[int] = []
    for fila in range(inicio, fin):
        rotulo = _texto(hoja.cell(row=fila, column=COL_B).value)
        if "denominaci" in rotulo and "estrategia" in rotulo:
            estrategia = fila
        elif rotulo.startswith(("área responsable", "area responsable")):
            area = fila
        elif rotulo.startswith("nombre del indicador"):
            indicadores.append(fila)
        elif rotulo.startswith("objetivo al que atiende"):
            objetivo = fila
        elif "descripci" in rotulo and "alcance" in rotulo:
            alcance = fila
        if _texto(hoja.cell(row=fila, column=6).value) == "prog":
            calendario = fila
    if None in (estrategia, area, objetivo, alcance, calendario) or not indicadores:
        return None
    assert estrategia is not None and area is not None and objetivo is not None
    assert alcance is not None and calendario is not None
    return Bloque(
        hoja=nombre,
        inicio=inicio,
        estrategia=estrategia,
        area=area,
        indicadores=indicadores,
        objetivo=objetivo,
        alcance=alcance,
        # Entre el encabezado del calendario y el alcance: dos filas de rótulos
        # (`PROG`/`ALC` y `#`/`%`) y después los huecos de actividad.
        actividades_desde=calendario + 2,
        actividades_hasta=alcance - 1,
    )


def limpiar_bloque(hoja, bloque: Bloque) -> None:
    """Deja el bloque en blanco, conservando rótulos y formato.

    Hace falta antes de escribir: un bloque que no recibe cédula se quedaría
    con los datos del año anterior, y nadie que abriera el archivo sabría que
    esa parte no se generó.
    """
    hoja.cell(row=bloque.estrategia + 2, column=COL_B).value = None
    hoja.cell(row=bloque.estrategia + 2, column=COL_P).value = None
    hoja.cell(row=bloque.estrategia + 2, column=COL_X).value = None
    hoja.cell(row=bloque.area + 2, column=COL_B).value = None
    hoja.cell(row=bloque.objetivo, column=COL_C).value = None
    hoja.cell(row=bloque.alcance + 1, column=COL_B).value = None
    for columna in (COL_B, COL_P):
        hoja.cell(row=bloque.alcance + 4, column=columna).value = None
        hoja.cell(row=bloque.alcance + 5, column=columna).value = None
    for fila in bloque.indicadores:
        for columna in (COL_B, COL_C, COL_D, COL_E, COL_J, COL_P):
            hoja.cell(row=fila + 1, column=columna).value = None
        for columna in (COL_J, COL_M, COL_P, COL_S, COL_V, COL_Y):
            hoja.cell(row=fila + 3, column=columna).value = None
    for fila in range(bloque.actividades_desde, bloque.actividades_hasta + 1):
        for columna in (COL_B, COL_D, COL_E, COL_R, COL_Y):
            hoja.cell(row=fila, column=columna).value = None
        for columnas in COLUMNAS_CUATRIMESTRE.values():
            for columna in columnas:
                hoja.cell(row=fila, column=columna).value = None


def _proporcion(cantidad: Decimal | None, meta: Decimal) -> float | None:
    """La columna `%` del calendario es la parte de la meta anual que toca a ese
    cuatrimestre, guardada como fracción porque la celda lleva formato de
    porcentaje: `0.33` se ve como 33 %. El sistema no la almacena —es
    derivada—, así que se calcula aquí."""
    if cantidad is None or meta <= 0:
        return None
    return float(cantidad) / float(meta)


def _numero(valor: Decimal | None) -> float | None:
    return None if valor is None else float(valor)


def escribir_bloque(hoja, bloque: Bloque, cedula: CedulaDelFormato) -> list[str]:
    """Escribe una cédula en su bloque y devuelve lo que no cupo."""
    incidencias: list[str] = []

    hoja.cell(row=bloque.estrategia + 2, column=COL_B).value = cedula.strategy_denomination
    hoja.cell(row=bloque.estrategia + 2, column=COL_P).value = cedula.objective_number
    hoja.cell(row=bloque.estrategia + 2, column=COL_X).value = cedula.strategy_type
    hoja.cell(row=bloque.area + 2, column=COL_B).value = cedula.responsible_area
    hoja.cell(row=bloque.objetivo, column=COL_C).value = cedula.objective_denomination
    hoja.cell(row=bloque.alcance + 1, column=COL_B).value = cedula.scope_and_socioeconomic_effect

    for indice, columna in enumerate((COL_B, COL_P)):
        if indice < len(cedula.signatories):
            nombre, cargo = cedula.signatories[indice]
            hoja.cell(row=bloque.alcance + 4, column=columna).value = nombre
            hoja.cell(row=bloque.alcance + 5, column=columna).value = cargo

    for indicador, fila in zip(cedula.indicators, bloque.indicadores, strict=False):
        hoja.cell(row=fila + 1, column=COL_B).value = indicador.name
        hoja.cell(row=fila + 1, column=COL_C).value = indicador.formula
        hoja.cell(row=fila + 1, column=COL_D).value = indicador.unit
        hoja.cell(row=fila + 1, column=COL_E).value = _numero(indicador.institutional_goal)
        hoja.cell(row=fila + 1, column=COL_J).value = indicador.baseline_year
        hoja.cell(row=fila + 3, column=COL_J).value = _numero(indicador.baseline_value)
        hoja.cell(row=fila + 3, column=COL_M).value = _porcentaje(indicador.current_percentage)
        hoja.cell(row=fila + 3, column=COL_P).value = _numero(indicador.target_value)
        hoja.cell(row=fila + 3, column=COL_S).value = _porcentaje(indicador.target_percentage)
        hoja.cell(row=fila + 3, column=COL_V).value = _numero(indicador.achieved_total)
        hoja.cell(row=fila + 3, column=COL_Y).value = _porcentaje(indicador.achieved_percentage)

    if len(cedula.indicators) > bloque.capacidad_indicadores:
        incidencias.append(
            f"La cédula de la estrategia {cedula.strategy_key} tiene "
            f"{len(cedula.indicators)} indicadores y el formato reserva "
            f"{bloque.capacidad_indicadores}: no se escribieron los últimos "
            f"{len(cedula.indicators) - bloque.capacidad_indicadores}."
        )

    filas = range(bloque.actividades_desde, bloque.actividades_hasta + 1)
    for actividad, fila in zip(cedula.activities, filas, strict=False):
        hoja.cell(row=fila, column=COL_B).value = con_clave(actividad.key, actividad.description)
        hoja.cell(row=fila, column=COL_D).value = actividad.unit
        hoja.cell(row=fila, column=COL_E).value = _numero(actividad.annual_goal)
        hoja.cell(row=fila, column=COL_R).value = actividad.deviation_justification
        hoja.cell(row=fila, column=COL_Y).value = actividad.executing_area
        for cuatrimestre, (prog, prog_pct, alc, alc_pct) in COLUMNAS_CUATRIMESTRE.items():
            programado = actividad.scheduled.get(cuatrimestre)
            alcanzado = actividad.achieved.get(cuatrimestre)
            hoja.cell(row=fila, column=prog).value = _numero(programado)
            hoja.cell(row=fila, column=prog_pct).value = _proporcion(
                programado, actividad.annual_goal
            )
            hoja.cell(row=fila, column=alc).value = _numero(alcanzado)
            hoja.cell(row=fila, column=alc_pct).value = _proporcion(
                alcanzado, actividad.annual_goal
            )

    if len(cedula.activities) > bloque.capacidad_actividades:
        sobrantes = [a.key for a in cedula.activities[bloque.capacidad_actividades :]]
        incidencias.append(
            f"La cédula de la estrategia {cedula.strategy_key} tiene "
            f"{len(cedula.activities)} actividades y el formato reserva "
            f"{bloque.capacidad_actividades}: quedaron fuera {', '.join(sobrantes)}."
        )

    return incidencias


def _porcentaje(valor: Decimal | None) -> float | None:
    """El sistema guarda los porcentajes de 0 a 100 y el formato los lleva como
    fracción, porque la celda tiene formato de porcentaje. Al importar se
    multiplicó por cien; aquí se deshace."""
    return None if valor is None else float(valor) / 100


def llenar_formato(
    plantilla: bytes, cedulas: list[CedulaDelFormato], *, anio: int
) -> tuple[bytes, list[str]]:
    """Rellena una copia de la plantilla y devuelve el archivo y lo que no cupo.

    Las cédulas se reparten por hoja de objetivo y, dentro de ella, en los
    huecos que el formato trae, en orden de clave de estrategia. Un objetivo
    con más cédulas que huecos no se recorta en silencio: se dice cuáles
    quedaron fuera.
    """
    from openpyxl import load_workbook

    libro = load_workbook(BytesIO(plantilla), keep_vba=True)
    bloques = localizar_bloques(libro)

    por_hoja: dict[str, list[Bloque]] = {}
    for bloque in bloques:
        por_hoja.setdefault(bloque.hoja, []).append(bloque)
        limpiar_bloque(libro[bloque.hoja], bloque)
        _poner_anio(libro[bloque.hoja], bloque, anio)

    incidencias: list[str] = []
    for numero, suyas in _por_objetivo(cedulas).items():
        huecos = _huecos_del_objetivo(por_hoja, numero)
        if not huecos:
            incidencias.append(
                f"El formato no trae hoja para el objetivo {numero}: quedaron fuera "
                f"{len(suyas)} cédulas ({', '.join(c.strategy_key for c in suyas)})."
            )
            continue
        if len(suyas) > len(huecos):
            fuera = [c.strategy_key for c in suyas[len(huecos) :]]
            incidencias.append(
                f"El objetivo {numero} tiene {len(suyas)} cédulas y el formato reserva "
                f"{len(huecos)} bloques: quedaron fuera {', '.join(fuera)}."
            )
        for cedula, bloque in zip(suyas, huecos, strict=False):
            incidencias.extend(escribir_bloque(libro[bloque.hoja], bloque, cedula))

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue(), incidencias


def _por_objetivo(cedulas: list[CedulaDelFormato]) -> dict[int, list[CedulaDelFormato]]:
    agrupadas: dict[int, list[CedulaDelFormato]] = {}
    for cedula in cedulas:
        agrupadas.setdefault(cedula.objective_number, []).append(cedula)
    for suyas in agrupadas.values():
        suyas.sort(key=lambda c: _orden_de_clave(c.strategy_key))
    return dict(sorted(agrupadas.items()))


def _orden_de_clave(clave: str) -> tuple[int, ...]:
    """`1.10` va después de `1.9`, que es lo que el orden alfabético invierte."""
    try:
        return tuple(int(parte) for parte in clave.split("."))
    except ValueError:
        return (0,)


def _huecos_del_objetivo(por_hoja: dict[str, list[Bloque]], numero: int) -> list[Bloque]:
    for nombre, bloques in por_hoja.items():
        encontrado = re.search(r"objetivo\s*(\d+)", _texto(nombre))
        if encontrado and int(encontrado.group(1)) == numero:
            return bloques
    return []


def _poner_anio(hoja, bloque: Bloque, anio: int) -> None:
    """El año aparece escrito en el título del bloque y en el encabezado del
    calendario. Se sustituye el que trae la plantilla para que el archivo no
    diga 2026 cuando se está entregando otro ejercicio."""
    for fila, columna in ((bloque.inicio, COL_B), (bloque.actividades_desde - 3, 6)):
        celda = hoja.cell(row=fila, column=columna)
        if isinstance(celda.value, str) and ANIO.search(celda.value):
            celda.value = ANIO.sub(str(anio), celda.value)
