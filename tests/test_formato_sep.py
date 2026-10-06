"""El POA de vuelta al formato de la SEP.

La prueba honesta de un exportador que rellena una plantilla es el viaje
redondo: escribir, volver a leer el archivo generado con el mismo mapa de
rótulos que usa el importador, y comprobar que sale lo que entró. Si difiere,
la exportación perdió algo por el camino y nadie se enteraría hasta que la SEP
devolviera el documento.
"""

from __future__ import annotations

import re
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.modules.poa_reports.application.sep_export import plantilla
from app.modules.poa_reports.application.sep_format import (
    ActividadDelFormato,
    CedulaDelFormato,
    IndicadorDelFormato,
    llenar_formato,
    localizar_bloques,
)

CLAVE_ACTIVIDAD = re.compile(r"^\s*(\d+\.\d+\.\d+)\b")


def _cedula(**cambios) -> CedulaDelFormato:
    base = {
        "objective_number": 1,
        "strategy_key": "1.1",
        "strategy_denomination": "Estrategia 1.1 Fortalecer las oportunidades educativas",
        "strategy_type": "Pertinencia",
        "responsible_area": "Secretaría Académica",
        "objective_denomination": "1. Aumentar la oferta educativa a nivel nacional",
        "scope_and_socioeconomic_effect": "Más jóvenes con acceso a la universidad.",
        "signatories": (
            ("Dra. Indra Ruth Toledo Coutiño", "Rectora"),
            ("Dr. Francisco Alberto Alonso Farrera", "Secretario Académico"),
        ),
        "indicators": (
            IndicadorDelFormato(
                name="1.1.15 Tasa neta de escolarización",
                formula="(a / b) * 100",
                unit="Porcentaje",
                institutional_goal=Decimal("68"),
                baseline_year=2025,
                baseline_value=Decimal("2117"),
                current_percentage=Decimal("68"),
                target_value=Decimal("2198"),
                target_percentage=Decimal("68"),
            ),
        ),
        "activities": (
            ActividadDelFormato(
                key="1.1.6",
                description="Desarrollar acciones gratuitas de formación",
                unit="Curso",
                annual_goal=Decimal("3"),
                executing_area="Dirección de Servicios Académicos",
                deviation_justification="2° cuatrimestre: se reprogramó un taller.",
                scheduled={1: Decimal("1"), 2: Decimal("1"), 3: Decimal("1")},
                achieved={1: Decimal("1"), 2: Decimal("0")},
            ),
        ),
    }
    base.update(cambios)
    return CedulaDelFormato(**base)


@pytest.fixture(scope="module")
def libro_de_una_cedula():
    """Rellenar la plantilla cuesta unos segundos —son 4,6 MiB con macros y mil
    quinientas celdas combinadas—, así que el caso común se arma una vez y las
    pruebas que sólo leen lo comparten."""
    archivo, incidencias = llenar_formato(plantilla(), [_cedula()], anio=2027)
    return load_workbook(BytesIO(archivo), keep_vba=True), incidencias


def _bloque_de(libro, clave: str):
    """El bloque donde quedó la cédula de esa estrategia."""
    for bloque in localizar_bloques(libro):
        hoja = libro[bloque.hoja]
        denominacion = hoja.cell(row=bloque.estrategia + 2, column=2).value
        if isinstance(denominacion, str) and f"Estrategia {clave} " in denominacion:
            return hoja, bloque
    raise AssertionError(f"No se escribió la estrategia {clave} en ningún bloque.")


def test_el_viaje_redondo_conserva_la_cedula(libro_de_una_cedula):
    libro, incidencias = libro_de_una_cedula

    assert incidencias == []
    hoja, bloque = _bloque_de(libro, "1.1")

    assert hoja.cell(row=bloque.estrategia + 2, column=2).value.startswith("Estrategia 1.1")
    assert hoja.cell(row=bloque.estrategia + 2, column=16).value == 1
    assert hoja.cell(row=bloque.estrategia + 2, column=24).value == "Pertinencia"
    assert hoja.cell(row=bloque.area + 2, column=2).value == "Secretaría Académica"
    assert hoja.cell(row=bloque.objetivo, column=3).value.startswith("1. Aumentar")
    assert hoja.cell(row=bloque.alcance + 1, column=2).value.startswith("Más jóvenes")

    # Los dos firmantes, cada uno con su cargo debajo.
    assert hoja.cell(row=bloque.alcance + 4, column=2).value == "Dra. Indra Ruth Toledo Coutiño"
    assert hoja.cell(row=bloque.alcance + 5, column=2).value == "Rectora"
    assert hoja.cell(row=bloque.alcance + 4, column=16).value.startswith("Dr. Francisco")
    assert hoja.cell(row=bloque.alcance + 5, column=16).value == "Secretario Académico"

    primero = bloque.indicadores[0]
    assert hoja.cell(row=primero + 1, column=2).value.startswith("1.1.15")
    assert hoja.cell(row=primero + 1, column=3).value == "(a / b) * 100"
    assert hoja.cell(row=primero + 1, column=4).value == "Porcentaje"
    assert hoja.cell(row=primero + 1, column=5).value == 68
    assert hoja.cell(row=primero + 1, column=10).value == 2025
    assert hoja.cell(row=primero + 3, column=10).value == 2117
    assert hoja.cell(row=primero + 3, column=16).value == 2198


def test_el_calendario_lleva_el_numero_y_su_parte_de_la_meta(libro_de_una_cedula):
    # La columna `%` del formato es la parte de la meta anual que toca a ese
    # cuatrimestre, guardada como fracción porque la celda tiene formato de
    # porcentaje. El sistema no la almacena: se calcula al escribir.
    libro, _ = libro_de_una_cedula
    hoja, bloque = _bloque_de(libro, "1.1")
    fila = bloque.actividades_desde

    assert hoja.cell(row=fila, column=2).value.startswith("1.1.6 Desarrollar")
    assert hoja.cell(row=fila, column=4).value == "Curso"
    assert hoja.cell(row=fila, column=5).value == 3
    assert hoja.cell(row=fila, column=6).value == 1  # 1er cuatrimestre, programado
    assert hoja.cell(row=fila, column=7).value == pytest.approx(1 / 3)
    assert hoja.cell(row=fila, column=8).value == 1  # alcanzado
    assert hoja.cell(row=fila, column=9).value == pytest.approx(1 / 3)
    assert hoja.cell(row=fila, column=14).value == 1  # 3er cuatrimestre, programado
    assert hoja.cell(row=fila, column=16).value is None  # sin alcanzado: no es cero
    assert hoja.cell(row=fila, column=25).value == "Dirección de Servicios Académicos"
    assert "se reprogramó" in hoja.cell(row=fila, column=18).value


def test_un_cuatrimestre_sin_reportar_queda_en_blanco_y_no_en_cero():
    # Escribir 0 diría que el área trabajó y no logró nada. Lo que pasa es que
    # todavía no ha reportado, y el formato lo dice dejando la celda vacía.
    cedula = _cedula(
        activities=(
            ActividadDelFormato(
                key="1.1.6",
                description="Una acción",
                unit="Informe",
                annual_goal=Decimal("3"),
                scheduled={1: Decimal("1")},
                achieved={},
            ),
        )
    )
    archivo, _ = llenar_formato(plantilla(), [cedula], anio=2027)
    libro = load_workbook(BytesIO(archivo), keep_vba=True)
    hoja, bloque = _bloque_de(libro, "1.1")
    fila = bloque.actividades_desde

    assert hoja.cell(row=fila, column=6).value == 1
    assert hoja.cell(row=fila, column=8).value is None
    assert hoja.cell(row=fila, column=9).value is None


def test_la_plantilla_va_en_blanco_y_lo_no_escrito_se_queda_en_blanco(libro_de_una_cedula):
    """Un bloque sin cédula no puede quedarse con los datos del año anterior.

    Si se quedara, quien abriera el archivo vería una cédula completa en una
    estrategia que este año no se formuló, y no tendría forma de saberlo.
    """
    libro, _ = libro_de_una_cedula

    escritos = 0
    for bloque in localizar_bloques(libro):
        hoja = libro[bloque.hoja]
        if hoja.cell(row=bloque.estrategia + 2, column=2).value:
            escritos += 1
            continue
        for fila in range(bloque.actividades_desde, bloque.actividades_hasta + 1):
            assert hoja.cell(row=fila, column=2).value is None
        assert hoja.cell(row=bloque.area + 2, column=2).value is None
    assert escritos == 1


def test_el_archivo_conserva_las_macros_y_el_formato(libro_de_una_cedula):
    """Rellenar, no generar. El .xlsm trae macros, celdas combinadas y el
    formato oficial: si al guardar se perdieran, lo que sale ya no sería el
    formato de la SEP aunque los datos estuvieran bien."""
    libro, _ = libro_de_una_cedula

    assert libro.vba_archive is not None
    assert len(libro.sheetnames) == 10
    hoja, _ = _bloque_de(libro, "1.1")
    assert len(list(hoja.merged_cells.ranges)) > 1000


def test_el_anio_del_ejercicio_reemplaza_al_de_la_plantilla():
    archivo, _ = llenar_formato(plantilla(), [_cedula()], anio=2029)
    libro = load_workbook(BytesIO(archivo), keep_vba=True)
    hoja, bloque = _bloque_de(libro, "1.1")

    assert "2029" in str(hoja.cell(row=bloque.inicio, column=2).value)
    assert "2026" not in str(hoja.cell(row=bloque.inicio, column=2).value)


def test_lo_que_no_cabe_se_dice_en_vez_de_recortarse():
    """El formato reserva cinco actividades por bloque y la plantilla no se
    puede estirar sin romper las celdas combinadas y el área de impresión. Una
    cédula más grande se escribe hasta donde cabe y lo demás se reporta."""
    actividades = tuple(
        ActividadDelFormato(
            key=f"1.1.{numero}",
            description=f"Acción {numero}",
            unit="Informe",
            annual_goal=Decimal("1"),
            scheduled={1: Decimal("1")},
            achieved={},
        )
        for numero in range(1, 8)
    )
    archivo, incidencias = llenar_formato(plantilla(), [_cedula(activities=actividades)], anio=2027)

    assert len(incidencias) == 1
    assert "1.1.6" in incidencias[0] and "1.1.7" in incidencias[0]
    libro = load_workbook(BytesIO(archivo), keep_vba=True)
    hoja, bloque = _bloque_de(libro, "1.1")
    escritas = [
        CLAVE_ACTIVIDAD.match(str(hoja.cell(row=fila, column=2).value or "")).group(1)
        for fila in range(bloque.actividades_desde, bloque.actividades_hasta + 1)
        if CLAVE_ACTIVIDAD.match(str(hoja.cell(row=fila, column=2).value or ""))
    ]
    assert escritas == ["1.1.1", "1.1.2", "1.1.3", "1.1.4", "1.1.5"]


def test_un_objetivo_con_mas_cedulas_que_bloques_lo_dice():
    # El objetivo 6 trae un solo bloque en la plantilla.
    cedulas = [
        _cedula(
            objective_number=6,
            strategy_key=f"6.{numero}",
            strategy_denomination=f"Estrategia 6.{numero} Una estrategia",
        )
        for numero in (1, 2, 3)
    ]

    _, incidencias = llenar_formato(plantilla(), cedulas, anio=2027)

    assert len(incidencias) == 1
    assert "objetivo 6" in incidencias[0]
    assert "6.2" in incidencias[0] and "6.3" in incidencias[0]


def test_las_estrategias_se_acomodan_por_su_numero_y_no_por_su_texto():
    # `1.10` va después de `1.9`: ordenar como texto las invertiría, y la
    # cédula acabaría en el bloque de otra estrategia.
    cedulas = [
        _cedula(strategy_key="1.10", strategy_denomination="Estrategia 1.10 La décima"),
        _cedula(strategy_key="1.9", strategy_denomination="Estrategia 1.9 La novena"),
    ]

    archivo, _ = llenar_formato(plantilla(), cedulas, anio=2027)

    libro = load_workbook(BytesIO(archivo), keep_vba=True)
    bloques = [b for b in localizar_bloques(libro) if "Objetivo 1" in b.hoja]
    escritas = [
        libro[b.hoja].cell(row=b.estrategia + 2, column=2).value
        for b in bloques
        if libro[b.hoja].cell(row=b.estrategia + 2, column=2).value
    ]
    assert escritas[0].startswith("Estrategia 1.9")
    assert escritas[1].startswith("Estrategia 1.10")


def test_rellenar_sobre_un_formato_ya_lleno_no_deja_restos_del_anterior():
    """Limpiar antes de escribir, y comprobarlo donde de verdad importa.

    La plantilla del repositorio va en blanco, así que limpiarla no se nota: una
    prueba que sólo la rellene pasaría igual si el limpiado desapareciera, y
    sobrescribir un bloque con otra cédula del mismo tamaño también lo tapa.

    Lo que de verdad protege es el caso de la cédula que **encoge**: si el año
    pasado tenía tres actividades y este tiene una, las otras dos se quedarían
    escritas, y quien abriera el archivo vería dos actividades que este
    ejercicio no se formularon.
    """

    def accion(numero: int) -> ActividadDelFormato:
        return ActividadDelFormato(
            key=f"1.1.{numero}",
            description=f"Acción {numero}",
            unit="Informe",
            annual_goal=Decimal("3"),
            scheduled={1: Decimal("1")},
            achieved={1: Decimal("1")},
        )

    llena, _ = llenar_formato(
        plantilla(), [_cedula(activities=(accion(1), accion(2), accion(3)))], anio=2027
    )
    encogida, _ = llenar_formato(llena, [_cedula(activities=(accion(1),))], anio=2027)

    libro = load_workbook(BytesIO(encogida), keep_vba=True)
    hoja, bloque = _bloque_de(libro, "1.1")

    assert hoja.cell(row=bloque.actividades_desde, column=2).value.startswith("1.1.1")
    sobrantes = [
        hoja.cell(row=fila, column=columna).value
        for fila in range(bloque.actividades_desde + 1, bloque.actividades_hasta + 1)
        for columna in (2, 4, 5, 6, 7, 8, 9)
    ]
    assert sobrantes == [None] * len(sobrantes)


def test_la_clave_no_se_repite_cuando_el_catalogo_ya_la_trae():
    """El catálogo del PSE guarda la descripción ya encabezada por su clave,
    porque así viene en el documento federal. Anteponerla sin mirar producía
    «1.1.6 1.1.6 Desarrollar…» en el formato que se le manda a la SEP."""
    cedula = _cedula(
        activities=(
            ActividadDelFormato(
                key="1.1.6",
                description="1.1.6 Desarrollar acciones gratuitas",
                unit="Curso",
                annual_goal=Decimal("3"),
                scheduled={},
                achieved={},
            ),
            ActividadDelFormato(
                key="1.1.9",
                description="Fortalecer los servicios",
                unit="Informe",
                annual_goal=Decimal("3"),
                scheduled={},
                achieved={},
            ),
        )
    )

    archivo, _ = llenar_formato(plantilla(), [cedula], anio=2027)

    libro = load_workbook(BytesIO(archivo), keep_vba=True)
    hoja, bloque = _bloque_de(libro, "1.1")
    assert hoja.cell(row=bloque.actividades_desde, column=2).value == (
        "1.1.6 Desarrollar acciones gratuitas"
    )
    # Y cuando falta, se antepone: el formato identifica la acción por su clave.
    assert hoja.cell(row=bloque.actividades_desde + 1, column=2).value == (
        "1.1.9 Fortalecer los servicios"
    )
